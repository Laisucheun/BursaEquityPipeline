"""One analyst workbook per company.

Sheets: Income Statement, Balance Sheet, Cash Flow (plus Changes in Equity when
any EQ facts exist), Ratios, Sources. Statement sheets share one layout:

    A: line item   B: concept key   C: unit   D..: fiscal years ascending

Amounts are divided by ``scale`` (default 1000 -> RM '000). Per-share concepts
(EPS in sen, NTA per share, ...) are written as stored, never rescaled.
"""

from __future__ import annotations

from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from bursa.analysis.dupont import compute_dupont
from bursa.analysis.facts import load_annual_facts
from bursa.analysis.growth import compute_growth
from bursa.db.models import Company
from bursa.export._common import annual_sources, concept_catalog, concept_info, concept_sort_key
from bursa.valuation.metrics import compute_valuation

STATEMENT_SHEETS = [
    ("IS", "Income Statement"),
    ("BS", "Balance Sheet"),
    ("CF", "Cash Flow"),
]
EQUITY_SHEET = ("EQ", "Changes in Equity")

HEADER_ROW = 1
FIRST_YEAR_COL = 4  # column D

AMOUNT_FMT = "#,##0;(#,##0);-"
AMOUNT_FMT_DEC = "#,##0.00;(#,##0.00);-"
PER_SHARE_FMT = "#,##0.00;(#,##0.00);-"
PCT_FMT = "0.0%"
TIMES_FMT = '0.00"x"'

BOLD = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="DDE4EE")


def unit_label(scale: int) -> str:
    return {1: "RM", 1000: "RM '000", 1_000_000: "RM mil"}.get(scale, f"RM / {scale:,}")


def _per_share_unit(key: str) -> str:
    return "sen" if key.startswith("is.eps") or key.endswith("_sen") else "per share"


def _header(ws: Worksheet, first: list[str], years: list[int]) -> None:
    for col, text in enumerate(first + [str(y) for y in years], start=1):
        cell = ws.cell(row=HEADER_ROW, column=col, value=text if col <= len(first) else int(text))
        cell.font = BOLD
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center" if col > len(first) else "left")


def _widths(ws: Worksheet, widths: dict[int, float], n_years: int, year_width: float = 14) -> None:
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    for i in range(n_years):
        ws.column_dimensions[get_column_letter(FIRST_YEAR_COL + i)].width = year_width


def _statement_sheet(ws, statement, annual, catalog, years, scale) -> None:
    keys = {k for f in annual.values() for k in f.values if concept_info(catalog, k).statement == statement}
    infos = sorted((concept_info(catalog, k) for k in keys), key=concept_sort_key)
    amount_fmt = AMOUNT_FMT if scale >= 1000 else AMOUNT_FMT_DEC

    _header(ws, ["Line item", "Concept", "Unit"], years)
    for row, info in enumerate(infos, start=HEADER_ROW + 1):
        unit = _per_share_unit(info.key) if info.is_per_share else unit_label(scale)
        label = info.label
        if info.is_per_share and f"({unit})" not in label.lower():
            label = f"{label} ({unit})"
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=info.key)
        ws.cell(row=row, column=3, value=unit)
        for i, fy in enumerate(years):
            value: Decimal | None = annual[fy].values.get(info.key)
            if value is None:
                continue
            out = float(value) if info.is_per_share else float(value / scale)
            cell = ws.cell(row=row, column=FIRST_YEAR_COL + i, value=out)
            cell.number_format = PER_SHARE_FMT if info.is_per_share else amount_fmt
        if info.is_subtotal:
            for col in range(1, FIRST_YEAR_COL + len(years)):
                ws.cell(row=row, column=col).font = BOLD

    ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=FIRST_YEAR_COL)
    _widths(ws, {1: 48, 2: 30, 3: 10}, len(years))


def _ratios_sheet(ws, session, company, years, scale) -> None:
    col_of = {fy: FIRST_YEAR_COL + i for i, fy in enumerate(years)}
    _header(ws, ["Metric", "Basis", "Unit"], years)
    row = HEADER_ROW + 1

    def section(title: str) -> None:
        nonlocal row
        row += 1
        ws.cell(row=row, column=1, value=title).font = BOLD
        row += 1

    def series(label, basis, unit, fmt, by_year: dict[int, float | None]) -> None:
        nonlocal row
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=basis)
        ws.cell(row=row, column=3, value=unit)
        for fy, v in by_year.items():
            if v is not None and fy in col_of:
                ws.cell(row=row, column=col_of[fy], value=float(v)).number_format = fmt
        row += 1

    dupont = compute_dupont(session, company)
    section("DuPont (3-factor)")
    d3 = {d.fiscal_year: d for d in dupont.three_factor}
    for attr, label, fmt in (("roe", "ROE", PCT_FMT), ("net_margin", "Net margin", PCT_FMT),
                             ("asset_turnover", "Asset turnover", TIMES_FMT),
                             ("equity_multiplier", "Equity multiplier", TIMES_FMT)):
        series(label, "DuPont 3", "%" if fmt == PCT_FMT else "x", fmt,
               {fy: getattr(d, attr) for fy, d in d3.items()})

    section("DuPont (5-factor)")
    d5 = {d.fiscal_year: d for d in dupont.five_factor}
    for attr, label, fmt in (("tax_burden", "Tax burden (PAT/PBT)", TIMES_FMT),
                             ("interest_burden", "Interest burden (PBT/EBIT)", TIMES_FMT),
                             ("operating_margin", "Operating margin (EBIT/Revenue)", PCT_FMT),
                             ("asset_turnover", "Asset turnover", TIMES_FMT),
                             ("equity_multiplier", "Equity multiplier", TIMES_FMT),
                             ("roe", "ROE", PCT_FMT)):
        series(label, "DuPont 5", "%" if fmt == PCT_FMT else "x", fmt,
               {fy: getattr(d, attr) for fy, d in d5.items()})

    section("Valuation inputs")
    val = {y.fiscal_year: y for y in compute_valuation(session, company).years}
    amount_fmt = AMOUNT_FMT if scale >= 1000 else AMOUNT_FMT_DEC
    for attr, label in (("ebit", "EBIT"), ("ebitda", "EBITDA"), ("nopat", "NOPAT"),
                        ("fcff", "FCFF"), ("fcfe", "FCFE"), ("net_debt", "Net debt"),
                        ("capex", "Capex"), ("dep_amort", "Depreciation & amortisation")):
        series(label, "valuation", unit_label(scale), amount_fmt,
               {fy: (None if getattr(y, attr) is None else getattr(y, attr) / scale)
                for fy, y in val.items()})
    series("Effective tax rate", "valuation", "%", PCT_FMT,
           {fy: y.effective_tax_rate for fy, y in val.items()})

    growth = compute_growth(session, company)
    section("ROE trend")
    series("ROE", "growth", "%", PCT_FMT, {r.fiscal_year: r.roe for r in growth.roe_trend})

    section("Growth (CAGR)")
    for col, text in enumerate(["Metric", "Span (yrs)", "Start FY", "End FY",
                                "Start value", "End value", "CAGR"], start=1):
        ws.cell(row=row, column=col, value=text).font = BOLD
    row += 1
    for m in (growth.revenue_cagr_3y, growth.revenue_cagr_5y, growth.earnings_cagr_3y,
              growth.earnings_cagr_5y, growth.asset_cagr_3y, growth.asset_cagr_5y):
        if m is None:
            continue
        ws.cell(row=row, column=1, value=f"{m.label} CAGR")
        ws.cell(row=row, column=2, value=m.years)
        ws.cell(row=row, column=3, value=m.start_year)
        ws.cell(row=row, column=4, value=m.end_year)
        ws.cell(row=row, column=5, value=m.start_value / scale).number_format = amount_fmt
        ws.cell(row=row, column=6, value=m.end_value / scale).number_format = amount_fmt
        ws.cell(row=row, column=7, value=m.cagr).number_format = PCT_FMT
        row += 1

    ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=FIRST_YEAR_COL)
    _widths(ws, {1: 36, 2: 12, 3: 10}, max(len(years), 4))


def _sources_sheet(ws, session, company, annual) -> None:
    sources = annual_sources(session, company)
    for col, text in enumerate(["Fiscal year", "Period end", "Concepts", "Source documents"], start=1):
        cell = ws.cell(row=HEADER_ROW, column=col, value=text)
        cell.font = BOLD
        cell.fill = HEADER_FILL
    for row, (fy, f) in enumerate(annual.items(), start=HEADER_ROW + 1):
        ws.cell(row=row, column=1, value=fy)
        ws.cell(row=row, column=2, value=f.period_end).number_format = "yyyy-mm-dd"
        ws.cell(row=row, column=3, value=len(f.values))
        ws.cell(row=row, column=4, value="; ".join(sources.get(fy, [])))
    ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=1)
    for col, w in {1: 12, 2: 12, 3: 10, 4: 90}.items():
        ws.column_dimensions[get_column_letter(col)].width = w


def company_workbook(session: Session, company: Company, *, scale: int = 1000) -> Workbook:
    if scale <= 0:
        raise ValueError("scale must be positive")
    annual = load_annual_facts(session, company)
    catalog = concept_catalog(session)
    years = sorted(annual)

    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = f"{company.stock_code} {company.name}"

    sheets = list(STATEMENT_SHEETS)
    if any(concept_info(catalog, k).statement == "EQ" for f in annual.values() for k in f.values):
        sheets.append(EQUITY_SHEET)
    for statement, title in sheets:
        _statement_sheet(wb.create_sheet(title), statement, annual, catalog, years, scale)
    _ratios_sheet(wb.create_sheet("Ratios"), session, company, years, scale)
    _sources_sheet(wb.create_sheet("Sources"), session, company, annual)
    return wb
